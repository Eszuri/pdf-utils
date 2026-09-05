"""
PDF to DOCX Standalone Converter CLI for PDF Utils
Usage: converter_cli.py <input_pdf_path> <output_docx_path>
"""

import sys
import os
import io

def detect_h_rules(page):
    """Detect groups of horizontal lines (decorative borders) in a PDF page."""
    pw = page.rect.width
    h_lines = []
    for d in page.get_cdrawings():
        lw = d.get('width', 0)
        rect = d.get('rect')
        if rect:
            x0, y0, x1, y1 = rect
            rw = abs(x1 - x0)
            rh = abs(y1 - y0)
            if rh < 15 and rw > pw * 0.25:
                h_lines.append({
                    'y': (y0 + y1) / 2,
                    'y0': min(y0, y1),
                    'y1': max(y0, y1),
                    'w': max(rh, lw, 0.5),
                    'x0': min(x0, x1),
                    'x1': max(x0, x1)
                })
                continue
        for item in d.get('items', []):
            if item[0] == 'l':
                p1, p2 = item[1], item[2]
                try:
                    dy, dx = abs(p1.y - p2.y), abs(p2.x - p1.x)
                except (AttributeError, TypeError):
                    continue
                if dy < 3 and dx > pw * 0.25:
                    h_lines.append({
                        'y': (p1.y + p2.y) / 2,
                        'y0': min(p1.y, p2.y) - max(lw, 0.5)/2,
                        'y1': max(p1.y, p2.y) + max(lw, 0.5)/2,
                        'w': max(lw, 0.5),
                        'x0': min(p1.x, p2.x),
                        'x1': max(p1.x, p2.x),
                    })
    if not h_lines:
        return []
    h_lines.sort(key=lambda x: x['y'])
    groups, cur = [], [h_lines[0]]
    for i in range(1, len(h_lines)):
        if h_lines[i]['y'] - cur[-1]['y'] < 20:
            cur.append(h_lines[i])
        else:
            groups.append(cur)
            cur = [h_lines[i]]
    groups.append(cur)
    return groups

def rasterize_rules(fitz, page, group, margin=2):
    """Rasterize a group of horizontal rules as a high-quality PNG image."""
    y_min = min(r['y0'] for r in group) - margin
    y_max = max(r['y1'] for r in group) + margin
    x_min = min(r['x0'] for r in group) - margin
    x_max = max(r['x1'] for r in group) + margin
    clip = fitz.Rect(
        max(x_min, page.rect.x0), max(y_min, page.rect.y0),
        min(x_max, page.rect.x1), min(y_max, page.rect.y1),
    )
    pix = page.get_pixmap(matrix=fitz.Matrix(4, 4), clip=clip, alpha=True)
    return pix.tobytes("png")

def find_anchor(page, group):
    """Find nearest text blocks above and below a horizontal rule group."""
    y_top = min(r['y0'] for r in group) - 2
    y_bot = max(r['y1'] for r in group) + 2
    below, below_d = None, float('inf')
    above, above_d = None, float('inf')
    for b in page.get_text("dict")["blocks"]:
        if b["type"] != 0:
            continue
        txt = "".join(
            s.get("text", "") for l in b.get("lines", []) for s in l.get("spans", [])
        ).strip()
        if len(txt) < 2:
            continue
        if b["bbox"][1] >= y_bot and b["bbox"][1] - y_bot < below_d:
            below, below_d = txt, b["bbox"][1] - y_bot
        if b["bbox"][3] <= y_top and y_top - b["bbox"][3] < above_d:
            above, above_d = txt, y_top - b["bbox"][3]
    return below, above

def insert_rule_image(doc, anchor_para, img_bytes, position):
    """Insert a rasterized rule image before or after the anchor paragraph."""
    from docx.shared import Emu
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    img_stream = io.BytesIO(img_bytes)
    try:
        sec = doc.sections[0]
        avail = sec.page_width - sec.left_margin - sec.right_margin
    except Exception:
        avail = Emu(int(6.0 * 914400))
    doc.add_picture(img_stream, width=avail)
    pic_elem = doc.element.body[-1]
    pPr = pic_elem.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        pic_elem.insert(0, pPr)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    pPr.append(jc)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:after'), '60')
    spacing.set(qn('w:line'), '240')
    pPr.append(spacing)
    if position == 'before':
        anchor_para._element.addprevious(pic_elem)
    else:
        anchor_para._element.addnext(pic_elem)

def postprocess(inp, out):
    """Detect horizontal rules in the PDF and insert them as images in the DOCX."""
    try:
        import fitz
        from docx import Document
    except ImportError:
        return

    try:
        pdf = fitz.open(inp)
        doc = Document(out)
        changed = False
        sf = 0
        paras = doc.paragraphs
        for pi in range(len(pdf)):
            for group in detect_h_rules(pdf[pi]):
                img_bytes = rasterize_rules(fitz, pdf[pi], group)
                below, above = find_anchor(pdf[pi], group)
                placed = False
                if below:
                    needle = ' '.join(below.split()[:4])
                    for i in range(sf, len(paras)):
                        if needle.lower() in paras[i].text.lower():
                            insert_rule_image(doc, paras[i], img_bytes, 'before')
                            sf = i + 1
                            changed = placed = True
                            break
                if not placed and above:
                    needle = ' '.join(above.split()[:4])
                    for i in range(sf, len(paras)):
                        if needle.lower() in paras[i].text.lower():
                            insert_rule_image(doc, paras[i], img_bytes, 'after')
                            sf = i + 1
                            changed = True
                            break
        if changed:
            doc.save(out)
        pdf.close()
    except Exception:
        pass

def convert_pdf_to_docx(inp, out):
    if not os.path.exists(inp):
        sys.stderr.write(f"Error: Input file not found: {inp}\n")
        sys.exit(1)

    try:
        from pdf2docx import Converter
    except ImportError:
        sys.stderr.write("Error: 'pdf2docx' is not installed.\n")
        sys.exit(2)

    try:
        cv = Converter(inp)
        cv.convert(
            out,
            shape_min_dimension=0.5,
            min_svg_h=0.5,
            min_svg_w=0.5,
            connected_border_tolerance=1.0,
            max_border_width=10.0,
            min_border_clearance=1.0,
            line_separate_threshold=2.0,
            line_overlap_threshold=0.95,
            lines_left_aligned_threshold=2.0,
            clip_image_res_ratio=6.0,
            new_paragraph_free_space_ratio=0.7,
            line_break_width_ratio=0.8,
            line_break_free_space_ratio=0.05,
        )
        cv.close()
    except Exception as e:
        sys.stderr.write(f"Conversion error: {e}\n")
        sys.exit(3)

    try:
        postprocess(inp, out)
    except Exception:
        pass

    if not os.path.exists(out):
        sys.stderr.write(f"Error: Output file was not generated: {out}\n")
        sys.exit(4)

    sys.stdout.write("Success\n")
    sys.exit(0)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.stderr.write("Usage: converter_cli.py <input.pdf> <output.docx>\n")
        sys.exit(1)

    convert_pdf_to_docx(sys.argv[1], sys.argv[2])
