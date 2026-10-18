"""Emit editable paragraphs and native Word tables using source geometry."""
import io
import re
import unicodedata
from statistics import median

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor
from .layout import reading_order
from .model import Line
from .cells import baseline, span_gaps, vertical_alignment, top_baselines


def clean_text(text):
    return ''.join(c for c in text if c in '\t\n\r' or '\x20' <= c <= '\ud7ff' or '\ue000' <= c <= '\ufffd' or '\U00010000' <= c <= '\U0010ffff')


def element(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def twips(value):
    return str(round(value * 20))


def font_family(name):
    # Remove PDF subset and style suffixes; formatting is emitted separately.
    name = re.sub(r'^[A-Z]{6}\+', '', name)
    name = re.sub(r'[-,](?:BoldItalic|BoldOblique|Bold|Italic|Oblique|Regular|Roman)(?:MT)?$', '', name, flags=re.I)
    return {'Helvetica': 'Arial', 'Times': 'Times New Roman', 'Courier': 'Courier New',
            'ArialMT': 'Arial', 'TimesNewRomanPSMT': 'Times New Roman',
            'TimesNewRomanPS': 'Times New Roman', 'SegoeUISymbol': 'Segoe UI Symbol',
            'SegoeUIEmoji': 'Segoe UI Emoji', 'SymbolMT': 'Segoe UI Symbol', 'Symbol': 'Segoe UI Symbol',
            'MalgunGothic': 'Malgun Gothic', 'MicrosoftJhengHeiRegular': 'Microsoft JhengHei'}.get(name, name)


def write_spans(paragraph, spans, tab_origin=None):
    content_size = max((s.size for s in spans if s.text.strip()), default=1)
    gaps = span_gaps(spans) if tab_origin is not None else {}
    for index,span in enumerate(spans):
        if index in gaps:
            position = max(0,gaps[index]-tab_origin)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(position))
            paragraph.add_run('\t')
        run = paragraph.add_run(clean_text(span.text))
        run.font.name = font_family(span.font)
        size = max(1, span.size if span.text.strip() else min(span.size, content_size))
        # DOCX uses half-points. python-docx truncates fractional half-points,
        # so a PDF's 9.96 pt would otherwise become 9.5 instead of 10 pt.
        half_points = int(size*2+.5)
        run.font.size = Pt(half_points/2)
        run.bold = bool(span.flags & 16)
        run.italic = bool(span.flags & 2)
        run.font.cs_bold = run.bold
        run.font.cs_italic = run.italic
        run.underline = bool(span.char_flags & 2)
        run.font.strike = bool(span.char_flags & 1)
        run.font.color.rgb = RGBColor.from_string(f'{span.color:06X}')
        props = run._r.get_or_add_rPr()
        fonts = props.find(qn('w:rFonts'))
        for attr in ('eastAsia', 'cs'):
            fonts.set(qn('w:' + attr), font_family(span.font))
        props.append(element('szCs', val=half_points))
        if any(unicodedata.bidirectional(c) in ('R', 'AL') for c in span.text):
            props.append(element('rtl'))
        # Preserve baseline offsets, including super/subscript, from geometry.
        baseline = max(s.origin[1] for s in spans if s.size >= max(x.size for x in spans)*.85)
        offset = baseline - span.origin[1]
        if abs(offset) > .7:
            props.append(element('position', val=round(offset*2)))


def compact(paragraph, size=1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(size)
    fmt.widow_control = False
    fmt.keep_with_next = False
    fmt.keep_together = False


class Writer:
    def __init__(self):
        self.doc = Document()
        normal = self.doc.styles['Normal']
        normal.paragraph_format.space_after = Pt(0)
        normal.font.size = Pt(10)
        # Headings carry outline semantics only. Source runs supply all visual
        # formatting, including complex-script text; template themes must not
        # introduce another font, italic face, or synthetic bold.
        for level in range(1, 7):
            for name in (f'Heading {level}', f'Heading {level} Char'):
                style = self.doc.styles[name]
                if style.element.rPr is not None:
                    style.element.remove(style.element.rPr)
        self.count = 0

    def page(self, page):
        section = self.doc.sections[0] if self.count == 0 else self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width = Pt(page.width)
        section.page_height = Pt(page.height)
        section.orientation = WD_ORIENT.LANDSCAPE if page.width > page.height else WD_ORIENT.PORTRAIT
        # Absolute source positions are relative to the page, not its margins.
        section.left_margin = section.right_margin = Pt(0)
        section.top_margin = section.bottom_margin = Pt(0)
        section.header_distance = section.footer_distance = Pt(0)
        if self.count == 0:
            anchor = self.doc.add_paragraph()
        else:
            anchor = self.doc.add_paragraph()
        compact(anchor)
        bookmark_id = str(page.number)
        anchor._p.append(element('bookmarkStart', id=bookmark_id, name=f'pdf_page_{page.number}'))
        anchor._p.append(element('bookmarkEnd', id=bookmark_id))
        if page.background:
            self.picture(anchor, page.background)
        for picture in page.pictures:
            self.picture(anchor, picture)
        for region in reading_order(page.lines + page.tables, page.height):
            if isinstance(region, Line):
                self.line(region, page.width)
            else:
                self.table(region)
        self.count += 1

    def line(self, line, page_width):
        p = self.doc.add_paragraph()
        if line.heading_level is not None:
            p.style = f'Heading {line.heading_level}'
        size = max(s.size for s in line.spans)
        compact(p, max(size*1.15, line.bbox[3]-line.bbox[1]))
        x, y, right, bottom = line.bbox
        width = min(page_width-x, max(2, right-x+3))
        frame = element('framePr', w=twips(width), x=twips(max(0,x)), y=twips(max(0,y)),
                        hAnchor='page', vAnchor='page', wrap='none', hSpace='0', vSpace='0')
        p._p.get_or_add_pPr().append(frame)
        write_spans(p, line.spans)

    def picture(self, anchor, picture):
        x0,y0,x1,y1 = picture.bbox
        inline = anchor.add_run().add_picture(io.BytesIO(picture.data), width=Pt(x1-x0), height=Pt(y1-y0))._inline
        if picture.svg:
            # Keep a standard PNG blip for readers without SVG support.
            from lxml import etree
            package = self.doc.part.package
            svg_part = Part(package.next_partname('/word/media/vector%d.svg'),
                            'image/svg+xml',picture.svg,package)
            relation = self.doc.part.relate_to(svg_part,RT.IMAGE)
            blip = inline.xpath('.//a:blip')[0]
            extensions = etree.SubElement(blip,qn('a:extLst'))
            extension = etree.SubElement(extensions,qn('a:ext'),
                                         uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
            svg_ns = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
            svg_blip = etree.SubElement(extension,f'{{{svg_ns}}}svgBlip',nsmap={'asvg':svg_ns})
            svg_blip.set(qn('r:embed'),relation)
        # Convert inline DrawingML to a floating picture behind editable text.
        from lxml import etree
        ns = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
        floating = etree.Element(f'{{{ns}}}anchor', nsmap=inline.nsmap,
                                distT='0', distB='0', distL='0', distR='0', simplePos='0',
                                relativeHeight='0', behindDoc='1', locked='0', layoutInCell='1', allowOverlap='1')
        etree.SubElement(floating, f'{{{ns}}}simplePos', x='0', y='0')
        for axis, position in [('H',x0), ('V',y0)]:
            pos = etree.SubElement(floating, f'{{{ns}}}position{axis}', relativeFrom='page')
            etree.SubElement(pos, f'{{{ns}}}posOffset').text = str(round(position*12700))
        floating.append(inline.extent)
        etree.SubElement(floating, f'{{{ns}}}wrapNone')
        for child in list(inline):
            floating.append(child)
        inline.getparent().replace(inline, floating)

    def table(self, table):
        rows, cols = len(table.ys)-1, len(table.xs)-1
        t = self.doc.add_table(rows=rows, cols=cols)
        t.autofit = False
        for i, width in enumerate(b-a for a,b in zip(table.xs,table.xs[1:])):
            t.columns[i].width = Pt(width)
            for cell in t.columns[i].cells:
                cell.width = Pt(width)
        props = t._tbl.tblPr
        props.append(element('tblpPr', leftFromText='0', rightFromText='0', topFromText='0', bottomFromText='0',
                             vertAnchor='page', horzAnchor='page', tblpX=twips(table.bbox[0]), tblpY=twips(table.bbox[1])))
        props.append(element('tblOverlap', val='overlap'))
        margins = element('tblCellMar')
        for side in ('top','bottom','left','right'):
            margins.append(element(side, w='0', type='dxa'))
        props.append(margins)
        borders = element('tblBorders')
        for side in ('top','bottom','left','right','insideH','insideV'):
            borders.append(element(side, val='single', sz=max(2,round(table.border_width*8)), color=table.border_color))
        props.append(borders)
        for ri, row in enumerate(t.rows):
            row.height = Pt(table.ys[ri+1]-table.ys[ri])
            # Minimum heights allow edits without truncating the source text.
        reference_baselines=top_baselines(table)
        row_top=[0]*rows
        row_bottom=[0]*rows
        for source in sorted(table.cells, key=lambda c:(c.row,c.col)):
            cell = t.cell(source.row, source.col)
            if source.rows > 1 or source.cols > 1:
                cell = cell.merge(t.cell(source.row+source.rows-1, source.col+source.cols-1))
            alignment=vertical_alignment(source)
            cell._tc.get_or_add_tcPr().append(element('vAlign', val=alignment))
            if source.fill:
                cell._tc.get_or_add_tcPr().append(element('shd', fill=source.fill))
            lines = sorted(source.lines, key=lambda line:(baseline(line),line.bbox[0]))
            pitches = [baseline(b)-baseline(a) for a,b in zip(lines,lines[1:])
                       if baseline(b)-baseline(a) > .5]
            content_height=0
            for i,line in enumerate(lines):
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                if line.heading_level is not None:
                    p.style = f'Heading {line.heading_level}'
                size = max((s.size for s in line.spans if s.text.strip()), default=1)
                regular = [value for value in pitches if size <= value <= size*1.6]
                leading = median(regular) if regular else size*1.12
                pitch = baseline(lines[i+1])-baseline(line) if i+1 < len(lines) else leading
                # Large blank gaps belong after a paragraph. Making them the
                # line height would move its text baseline and inflate the row.
                height = max(size, pitch if pitch <= size*1.6 else leading)
                compact(p, height)
                if i+1 < len(lines):
                    p.paragraph_format.space_after = Pt(max(0,pitch-height))
                content_height+=height+(max(0,pitch-height) if i+1 < len(lines) else 0)
                # Preserve cell insets measured from the text position.
                left = max(0,line.spans[0].bbox[0]-source.bbox[0])
                p.paragraph_format.left_indent = Pt(left)
                gaps = span_gaps(line.spans)
                if gaps:
                    split = min(gaps)
                    marker = ''.join(s.text for s in line.spans[:split]).strip()
                    if re.fullmatch(r'(?:[-*+\u2022\u25e6\u25aa\u2013\u2014]|\(?\d+[.)]|[A-Za-z][.)])',marker):
                        text_left = gaps[split]-source.bbox[0]
                        p.paragraph_format.left_indent = Pt(text_left)
                        p.paragraph_format.first_line_indent = Pt(left-text_left)
                write_spans(p,line.spans,tab_origin=source.bbox[0])
            if lines and any(s.text.strip() for line in lines for s in line.spans):
                first=lines[0]
                span=max(first.spans,key=lambda s:s.size)
                reference=reference_baselines.get((span.font,round(span.size,1)))
                if reference is not None:
                    available=max(0,source.bbox[3]-source.bbox[1]-content_height)
                    offset=min(available,max(0,baseline(first)-source.bbox[1]-reference))
                    if alignment == 'center':
                        difference=2*offset-available
                        top,bottom=max(0,difference),max(0,-difference)
                    elif alignment == 'bottom':
                        top,bottom=0,available-offset
                    else:
                        top,bottom=offset,0
                    margins=element('tcMar')
                    margins.append(element('top',w=twips(top),type='dxa'))
                    margins.append(element('bottom',w=twips(bottom),type='dxa'))
                    cell._tc.get_or_add_tcPr().append(margins)
                    row_top[source.row]=max(row_top[source.row],top)
                    last_row=source.row+source.rows-1
                    row_bottom[last_row]=max(row_bottom[last_row],bottom)
            if not source.lines:
                compact(cell.paragraphs[0], 1)
        # Word adds cell padding to a minimum row height. Keep the source outer
        # height while still allowing rows to grow when the user edits content.
        for ri,row in enumerate(t.rows):
            row.height=Pt(max(1,table.ys[ri+1]-table.ys[ri]-row_top[ri]-row_bottom[ri]))

    def save(self, path):
        # Every floating table needs a following paragraph in Word.
        compact(self.doc.add_paragraph())
        self.doc.save(path)
