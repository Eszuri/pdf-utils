import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import pymupdf
from lxml import etree

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from pdf_word.convert import convert, validate
from pdf_word.layout import reading_order
from pdf_word.markdown import HeadingParser
from pdf_word.model import Line, Span
from pdf_word.write import font_family
from pdf_word.cells import merge_cell_lines


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root/'input.pdf'
        self.target = self.root/'output.docx'

    def pdf(self,text='Arbitrary customer data 9842',width=612,height=792):
        doc=pymupdf.open()
        page=doc.new_page(width=width,height=height)
        page.insert_text((40,50),text)
        doc.save(self.source)
        doc.close()

    def xml(self):
        with ZipFile(self.target) as package:
            return etree.fromstring(package.read('word/document.xml'))

    def test_form_labels_never_replace_content(self):
        self.pdf('SERVICE REPORT DATA USER CUSTOMER_9842')
        result=convert(self.source,self.target)
        text=''.join(self.xml().xpath('//*[local-name()="t"]/text()'))
        self.assertIn('CUSTOMER_9842',text)
        self.assertNotIn('STEMPEL',text)
        self.assertEqual(result['pages'],1)

    def test_failure_preserves_existing_destination(self):
        self.pdf()
        self.target.write_bytes(b'existing document')
        with patch('pdf_word.convert.validate',side_effect=ValueError('validation failed')):
            with self.assertRaisesRegex(ValueError,'validation failed'):
                convert(self.source,self.target)
        self.assertEqual(self.target.read_bytes(),b'existing document')
        self.assertEqual(list(self.root.glob('.pdf-utils-*')),[])

    def test_cancellation_preserves_destination_and_cleans_temporary(self):
        self.pdf()
        self.target.write_bytes(b'original')
        marker=self.root/'cancel'
        def cancel(event):
            marker.touch()
        with patch.dict(os.environ,{'PDF_WORD_CANCEL':str(marker)}):
            with self.assertRaisesRegex(InterruptedError,'CANCELLED'):
                convert(self.source,self.target,cancel)
        self.assertEqual(self.target.read_bytes(),b'original')
        self.assertEqual(list(self.root.glob('.pdf-utils-*')),[])

    def test_image_only_page_requires_ocr(self):
        doc=pymupdf.open()
        page=doc.new_page()
        pix=pymupdf.Pixmap(pymupdf.csRGB,pymupdf.IRect(0,0,10,10),False)
        pix.clear_with(255)
        page.insert_image(pymupdf.Rect(0,0,400,600),stream=pix.tobytes('png'))
        doc.save(self.source); doc.close()
        with self.assertRaisesRegex(ValueError,'OCR_REQUIRED'):
            convert(self.source,self.target)
        self.assertFalse(self.target.exists())

    def test_encrypted_pdf_is_rejected(self):
        doc=pymupdf.open(); doc.new_page()
        doc.save(self.source,encryption=pymupdf.PDF_ENCRYPT_AES_256,user_pw='secret',owner_pw='owner')
        doc.close()
        with self.assertRaisesRegex(ValueError,'PASSWORD_REQUIRED'):
            convert(self.source,self.target)

    def test_native_table_and_mixed_page_geometry(self):
        doc=pymupdf.open()
        page=doc.new_page(width=600,height=800)
        for x in [40,140,240]: page.draw_line((x,100),(x,180))
        for y in [100,140,180]: page.draw_line((40,y),(240,y))
        for x,y,text in [(45,120,'A1'),(145,120,'B1'),(45,160,'A2'),(145,160,'B2')]:
            page.insert_text((x,y),text)
        page=doc.new_page(width=800,height=600)
        page.insert_text((40,50),'Landscape')
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        self.assertEqual(result['tables'],1)
        self.assertEqual(len(self.xml().xpath('//*[local-name()="tbl"]')),1)
        sizes=self.xml().xpath('//*[local-name()="pgSz"]')
        ns='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        self.assertEqual([(int(x.get(ns+'w')),int(x.get(ns+'h'))) for x in sizes],[(12000,16000),(16000,12000)])

    def test_unicode_paths_and_json_progress(self):
        self.pdf()
        self.source=self.source.rename(self.root/'dokumen kutip \' dan spasi.pdf')
        self.target=self.root/'hasil spasi.docx'
        events=[]
        result=convert(self.source,self.target,events.append)
        self.assertEqual(events[-1]['kind'],'result')
        self.assertEqual(result['outputPath'],str(self.target.resolve()))

    def test_font_sizes_round_to_nearest_half_point_without_shrinking(self):
        doc=pymupdf.open(); page=doc.new_page()
        samples=[(9.96,10),(9.999999,10),(10,10),(11.96,12),(12,12),
                 (9.5,9.5),(6.48,6.5),(6.96,7),(8.04,8),(18.96,19),
                 (12.25,12.5),(10.7,10.5)]
        for i,(source,_) in enumerate(samples):
            page.insert_text((40,40+i*25),f'Size sample {i}',fontsize=source)
        doc.save(self.source); doc.close()
        convert(self.source,self.target)
        from docx import Document
        paragraphs={p.text:p for p in Document(self.target).paragraphs}
        for i,(_,expected) in enumerate(samples):
            with self.subTest(sample=i):
                run=paragraphs[f'Size sample {i}'].runs[0]
                self.assertEqual(run.font.size.pt,expected)
                size=run._r.xpath('./w:rPr/w:sz/@w:val')[0]
                complex_size=run._r.xpath('./w:rPr/w:szCs/@w:val')[0]
                self.assertEqual(size,complex_size)
                self.assertEqual(int(size),int(expected*2))

    def test_table_and_heading_use_the_same_font_size_quantization(self):
        doc=pymupdf.open(); page=doc.new_page()
        page.insert_text((40,50),'### Heading at ten',fontsize=9.96)
        for x in [40,180,320]: page.draw_line((x,100),(x,180))
        for y in [100,140,180]: page.draw_line((40,y),(320,y))
        page.insert_text((45,120),'Cell at ten',fontsize=9.96)
        page.insert_text((185,120),'Cell at nine and a half',fontsize=9.5)
        page.insert_text((45,160),'Cell at twelve',fontsize=12)
        doc.save(self.source); doc.close()
        convert(self.source,self.target)
        from docx import Document
        doc=Document(self.target)
        heading=next(p for p in doc.paragraphs if p.text=='Heading at ten')
        self.assertEqual(heading.runs[0].font.size.pt,10)
        table=doc.tables[0]
        for row,col,expected in [(0,0,10),(0,1,9.5),(1,0,12)]:
            self.assertEqual(table.cell(row,col).paragraphs[0].runs[0].font.size.pt,expected)

    def test_text_order_is_column_first(self):
        def line(x,y):
            box=(x,y,x+40,y+10)
            return Line([Span(str(x),box,(x,y+8),'Arial',10,0,0,0)],box)
        lines=[line(x,y) for y in range(100,220,20) for x in [40,240]]
        ordered=reading_order(lines,800)
        self.assertEqual([x.bbox[0] for x in ordered],[40]*6+[240]*6)

    def test_table_list_markers_and_items_share_paragraph_and_source_pitch(self):
        doc=pymupdf.open(); page=doc.new_page()
        for x in [40,240,440]: page.draw_line((x,80),(x,270))
        for y in [80,100,270]: page.draw_line((40,y),(440,y))
        page.insert_text((45,112),'List in cell')
        for i in range(10):
            page.insert_text((70,128+i*14),'-')
            if i < 9:
                page.insert_text((100,128+i*14),f'Item {i+1}')
        page.insert_text((245,128),'Other cell')
        page.insert_text((245,242),'Separated paragraph')
        page.insert_text((40,310),'Body below the table')
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        self.assertEqual(result['tables'],1)
        from docx import Document
        table=Document(self.target).tables[0]
        paragraphs=table.cell(1,0).paragraphs
        self.assertEqual(len(paragraphs),11)
        self.assertEqual(paragraphs[0].text,'List in cell')
        for i,p in enumerate(paragraphs[1:10],1):
            self.assertEqual(p.text,f'-\tItem {i}')
            self.assertAlmostEqual(p.paragraph_format.left_indent.pt,60,delta=.1)
            self.assertAlmostEqual(p.paragraph_format.first_line_indent.pt,-30,delta=.1)
            self.assertAlmostEqual(p.paragraph_format.line_spacing.pt,14,delta=.1)
        self.assertEqual(paragraphs[-1].text,'-')
        self.assertEqual(table.cell(1,1).text,'Other cell\nSeparated paragraph')
        spaced=table.cell(1,1).paragraphs[0].paragraph_format
        self.assertLess(spaced.line_spacing.pt,20)
        self.assertGreater(spaced.space_after.pt,90)

    def test_cell_line_merge_preserves_styles_and_wrapped_continuation(self):
        def fragment(text,x,y,size=11,flags=0):
            box=(x,y-size,x+len(text)*5,y+3)
            return Line([Span(text,box,(x,y),'Arial',size,flags,0,0)],box)
        # Font runs can differ slightly in baseline; PDF block order is irrelevant.
        lines=[fragment('Item',100,130,flags=16),fragment('-',70,130.2),
               fragment('continuation',100,144),fragment('1.',70,158),
               fragment('Second',100,158)]
        merged=merge_cell_lines(lines)
        self.assertEqual([''.join(s.text for s in line.spans) for line in merged],
                         ['-Item','continuation','1.Second'])
        self.assertEqual(merged[0].spans[1].flags,16)
        self.assertEqual(len(lines),5)

    def test_table_vertical_alignment_follows_source_geometry(self):
        doc=pymupdf.open(); page=doc.new_page()
        for x in [40,160,280,400]: page.draw_line((x,80),(x,260))
        for y in [80,140,260]: page.draw_line((40,y),(400,y))
        for x,y,text in [(45,92,'Top'),(165,112,'Center'),(285,136,'Bottom'),
                         (45,152,'Following row')]:
            page.insert_text((x,y),text)
        doc.save(self.source); doc.close()
        convert(self.source,self.target)
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        cells=self.xml().xpath('//w:tbl/w:tr[1]/w:tc',namespaces=ns)
        self.assertEqual([c.xpath('./w:tcPr/w:vAlign/@w:val',namespaces=ns)[0] for c in cells],
                         ['top','center','bottom'])
        self.assertTrue(cells[1].xpath('./w:tcPr/w:tcMar/w:bottom[@w:w>0]',namespaces=ns))
        self.assertEqual([len(c.findall('w:p',ns)) for c in cells],[1,1,1])

    def test_partial_underlines_and_cell_artwork_survive_border_filtering(self):
        doc=pymupdf.open(); page=doc.new_page()
        for x in [40,240,440]: page.draw_line((x,80),(x,250))
        for y in [80,140,250]: page.draw_line((40,y),(440,y))
        text='Plain Underlined Plain'
        start=45+pymupdf.get_text_length('Plain ',fontsize=12)
        end=start+pymupdf.get_text_length('Underlined',fontsize=12)
        for baseline_y in [158,178]:
            page.insert_text((45,baseline_y),text,fontsize=12)
        page.draw_line((start,160),(end,160),width=.6)
        page.draw_rect(pymupdf.Rect(start,180,end,180.7),color=None,fill=(0,0,0))
        page.draw_rect(pymupdf.Rect(70,195,110,213),color=(0,0,1))
        page.insert_text((45,248),'Border stays plain',fontsize=12)
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        from docx import Document
        paragraphs=Document(self.target).tables[0].cell(1,0).paragraphs
        for p in paragraphs[:2]:
            self.assertEqual(p.text,text)
            self.assertEqual(''.join(run.text for run in p.runs if run.underline),'Underlined')
        self.assertFalse(any(run.underline for run in paragraphs[-1].runs))
        self.assertEqual(result['graphicsPages'],1)
        with ZipFile(self.target) as package:
            svg=etree.fromstring(package.read(next(n for n in package.namelist() if n.endswith('.svg'))))
            self.assertTrue(svg.xpath('//*[local-name()="path"]'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        borders=self.xml().xpath('//w:tblBorders/w:top/@w:color',namespaces=ns)
        self.assertEqual(borders,['000000'])

    def test_vector_artwork_preserved_without_rasterizing_text(self):
        doc=pymupdf.open(); page=doc.new_page()
        page.insert_text((40,50),'Editable label')
        page.draw_rect(pymupdf.Rect(40,80,160,130),fill=(1,0,0),fill_opacity=.5)
        page.draw_bezier((40,200),(60,160),(100,250),(160,200),color=(0,0,1))
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        self.assertEqual(result['graphicsPages'],1)
        self.assertEqual(result['images'],0)
        self.assertIn('Editable label',''.join(self.xml().xpath('//*[local-name()="t"]/text()')))
        with ZipFile(self.target) as package:
            vectors=[name for name in package.namelist() if name.endswith('.svg')]
            self.assertEqual(len(vectors),1)
            svg=etree.fromstring(package.read(vectors[0]))
            self.assertTrue(svg.xpath('//*[local-name()="path"]'))
            self.assertFalse(svg.xpath('//*[local-name()="text"]'))
            self.assertTrue(self.xml().xpath('//*[local-name()="svgBlip"]'))

    def test_source_image_pixels_are_not_downsampled(self):
        doc=pymupdf.open(); page=doc.new_page()
        page.insert_text((40,50),'### Editable heading')
        pix=pymupdf.Pixmap(pymupdf.csRGB,pymupdf.IRect(0,0,1200,600),False)
        pix.clear_with(190)
        page.insert_image(pymupdf.Rect(40,80,160,140),stream=pix.tobytes('png'))
        doc.save(self.source); doc.close()
        with pymupdf.open(self.source) as pdf:
            original=next(b['image'] for b in pdf[0].get_text('dict')['blocks'] if b['type']==1)
        convert(self.source,self.target)
        with ZipFile(self.target) as package:
            name=next(name for name in package.namelist() if name.startswith('word/media/'))
            data=package.read(name)
            self.assertEqual(data,original)
            embedded=pymupdf.Pixmap(data)
            self.assertEqual((embedded.width,embedded.height),(1200,600))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        heading=self.xml().xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading3"]]',namespaces=ns)[0]
        self.assertEqual(heading.xpath('.//w:t/text()',namespaces=ns),['Editable heading'])
        self.assertFalse(heading.xpath('.//*[local-name()="drawing"]'))

    def test_image_mask_keeps_native_resolution_and_alpha(self):
        doc=pymupdf.open(); page=doc.new_page()
        page.insert_text((40,50),'Transparent image')
        pix=pymupdf.Pixmap(pymupdf.csRGB,pymupdf.IRect(0,0,120,60),True)
        pix.clear_with(200)
        pix.set_alpha(bytes([128]*(120*60)))
        page.insert_image(pymupdf.Rect(40,80,100,110),stream=pix.tobytes('png'))
        doc.save(self.source); doc.close()
        convert(self.source,self.target)
        with ZipFile(self.target) as package:
            name=next(name for name in package.namelist() if name.startswith('word/media/'))
            embedded=pymupdf.Pixmap(package.read(name))
            self.assertEqual((embedded.width,embedded.height),(120,60))
            self.assertTrue(embedded.alpha)

    def test_corrupt_pdf_does_not_commit_output(self):
        self.source.write_bytes(b'not a pdf')
        self.target.write_bytes(b'original')
        with self.assertRaises(Exception):
            convert(self.source,self.target)
        self.assertEqual(self.target.read_bytes(),b'original')
        self.assertEqual(list(self.root.glob('.pdf-utils-*')),[])

    def paragraphs(self):
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        return [(''.join(p.xpath('.//w:t/text()',namespaces=ns)),
                 p.xpath('./w:pPr/w:pStyle/@w:val',namespaces=ns))
                for p in self.xml().xpath('//w:p',namespaces=ns)]

    def test_markdown_heading_levels_and_closing_markers(self):
        doc=pymupdf.open(); page=doc.new_page()
        for level in range(1,7):
            page.insert_text((40,40+level*30),'#'*level+f' Title {level}')
        page.insert_text((40,260),'### Title #42 ###')
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        paragraphs=dict(self.paragraphs())
        for level in range(1,7):
            self.assertEqual(paragraphs[f'Title {level}'],[f'Heading{level}'])
        self.assertEqual(paragraphs['Title #42'],['Heading3'])
        self.assertEqual(result['markdownHeadings'],7)

    def test_literal_hashes_and_code_remain_unchanged(self):
        lines=['C# and tag #123', '###no-space', '####### Too many hashes',
               r'\### Escaped', '    ### Indented code', '###',
               '```python','### Fenced code','```',
               '~~~','## Tilde code','~~~']
        doc=pymupdf.open(); page=doc.new_page()
        for i,text in enumerate(lines):
            page.insert_text((40,40+i*20),text)
        page.insert_text((40,320),'### Monospace code',fontname='cour')
        page.insert_text((40,350),'### Actual heading')
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        paragraphs=dict(self.paragraphs())
        for text in lines+['### Monospace code']:
            self.assertIn(text,paragraphs)
            self.assertFalse(paragraphs[text])
        self.assertEqual(paragraphs['Actual heading'],['Heading3'])
        self.assertEqual(result['markdownHeadings'],1)

    def test_markdown_marker_across_spans_retains_title_format(self):
        def span(text,flags=0):
            return Span(text,(40,40,200,52),(40,50),'Arial',10,flags,0,0)
        original=Line([span('#'),span('## '),span('Bold title',16),span(' #12 ###')],(40,40,200,52))
        converted,removed=HeadingParser().parse(original)
        self.assertEqual(converted.heading_level,3)
        self.assertEqual(''.join(s.text for s in converted.spans),'Bold title #12')
        self.assertEqual(converted.spans[0].flags,16)
        self.assertEqual(removed,'###  ###')
        self.assertEqual(''.join(s.text for s in original.spans),'### Bold title #12 ###')

    def test_heading_typography_comes_from_source_not_template(self):
        doc=pymupdf.open(); page=doc.new_page()
        fonts=[('helv','Arial',False,False),('hebo','Arial',True,False),
               ('tiit','Times New Roman',False,True)]*2
        for i,(font,_,_,_) in enumerate(fonts,1):
            page.insert_text((40,40+i*40),'#'*i+f' Heading {i}',
                             fontname=font,fontsize=12+i,color=(0,0,0))
        doc.save(self.source); doc.close()
        convert(self.source,self.target)
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        w='{'+ns['w']+'}'
        for i,(_,family,bold,italic) in enumerate(fonts,1):
            heading=self.xml().xpath(f'//w:p[w:pPr/w:pStyle[@w:val="Heading{i}"]]',namespaces=ns)[0]
            self.assertEqual(heading.xpath('.//w:t/text()',namespaces=ns),[f'Heading {i}'])
            props=heading.find(w+'r/'+w+'rPr')
            self.assertEqual(props.find(w+'rFonts').get(w+'ascii'),family)
            self.assertEqual(props.find(w+'sz').get(w+'val'),str((12+i)*2))
            for tag,value in [('b',bold),('bCs',bold),('i',italic),('iCs',italic)]:
                self.assertEqual(props.find(w+tag).get(w+'val')!='0',value)
            self.assertFalse(heading.xpath('.//*[local-name()="drawing"]'))
        with ZipFile(self.target) as package:
            styles=etree.fromstring(package.read('word/styles.xml'))
            for i in range(1,7):
                for style in (f'Heading{i}',f'Heading{i}Char'):
                    self.assertFalse(styles.xpath(f'//w:style[@w:styleId="{style}"]/w:rPr',namespaces=ns))

    def test_postscript_font_names_resolve_to_word_families(self):
        for source,family in [('ABCDEF+Arial-BoldMT','Arial'),
                              ('Arial-BoldItalicMT','Arial'),
                              ('TimesNewRomanPS-BoldMT','Times New Roman'),
                              ('Helvetica-BoldOblique','Arial'),
                              ('Courier-Oblique','Courier New'),
                              ('CustomFont-Bold','CustomFont')]:
            with self.subTest(source=source):
                self.assertEqual(font_family(source),family)

    def test_heading_with_pdf_nonbreaking_spaces(self):
        text='###\u00a0Judul jelas\u00a0###'
        original=Line([Span(text,(40,40,200,58),(40,54),'Arial-BoldMT',14,16,0,0)],(40,40,200,58))
        converted,removed=HeadingParser().parse(original)
        self.assertEqual(converted.heading_level,3)
        self.assertEqual(''.join(s.text for s in converted.spans),'Judul jelas')
        self.assertEqual(removed,'###\u00a0\u00a0###')

    def test_fenced_code_state_survives_page_boundary(self):
        doc=pymupdf.open(); page=doc.new_page()
        page.insert_text((40,50),'```')
        page=doc.new_page()
        for i,text in enumerate(['### Still code','```','### Actual heading']):
            page.insert_text((40,50+i*25),text)
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        paragraphs=dict(self.paragraphs())
        self.assertFalse(paragraphs['### Still code'])
        self.assertEqual(paragraphs['Actual heading'],['Heading3'])
        self.assertEqual(result['markdownHeadings'],1)

    def test_headings_in_table_cells_have_independent_code_fences(self):
        doc=pymupdf.open(); page=doc.new_page()
        for x in [40,240,440]: page.draw_line((x,100),(x,240))
        for y in [100,180,240]: page.draw_line((40,y),(440,y))
        for x,y,text in [(45,120,'```'),(45,140,'### Code in cell'),
                         (245,120,'### Cell title'),(45,200,'A2'),(245,200,'B2')]:
            page.insert_text((x,y),text)
        page.insert_text((40,280),'### Body title')
        doc.save(self.source); doc.close()
        result=convert(self.source,self.target)
        paragraphs=dict(self.paragraphs())
        self.assertEqual(result['tables'],1)
        self.assertEqual(result['markdownHeadings'],2)
        self.assertFalse(paragraphs['### Code in cell'])
        self.assertEqual(paragraphs['Cell title'],['Heading3'])
        self.assertEqual(paragraphs['Body title'],['Heading3'])

    def test_validation_still_rejects_missing_heading_title(self):
        self.pdf('### Important title #42')
        convert(self.source,self.target)
        from docx import Document
        doc=Document(self.target)
        paragraph=next(p for p in doc.paragraphs if p.text=='Important title #42')
        paragraph.text='Important #42'
        doc.save(self.target)
        with self.assertRaisesRegex(ValueError,'TEXT_VALIDATION_FAILED'):
            validate(self.target,'### Important title #42',0,'### ')


if __name__=='__main__':
    unittest.main()
